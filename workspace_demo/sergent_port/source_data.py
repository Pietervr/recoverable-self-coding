#!/usr/bin/env python3
"""Reader for the publisher's Source Data workbook of Sergent et al. 2021
(41467_2021_21393_MOESM4_ESM.xlsx, Springer Nature; see README "Source Data provenance").

Blocks used by reconcile.py (coordinates verified with `python source_data.py --inspect`):
  * 'Figure 3 Panel E'              rows 2-4, cols A..BA : pxp of Null / unimodal 2B / bifurcation per
                                    window (53 windows; PlotFig convention t = time(TWOI(i)) + 15 ms)
  * 'Figure 1 Panel D'              rows 2-21, cols Q..V : SD of audibility (% scale) per subject, 6 levels
  * 'Figure 2 Panel C rightmost col' rows 26-45, cols A..AN : neural variability profile per subject,
                                    8 windows x 5 levels (-13 .. -5 dB)
  * 'Supplementary Figure 4'        rows 25-44, cols A..AJ : passive variability profile per subject,
                                    6 windows (0-100 .. 500-600 ms) x 6 levels
"""
from __future__ import annotations

import os
import sys

import numpy as np
import openpyxl

from common import DATA_ROOT

XLSX = os.path.join(DATA_ROOT, "41467_2021_21393_MOESM4_ESM.xlsx")
XLSX_URL = ("https://media.springernature.com/original/springer-static/esm/art%3A10.1038%2Fs41467-021-21393-z/"
            "MediaObjects/41467_2021_21393_MOESM4_ESM.xlsx")

_wb = None


def workbook():
    global _wb
    if _wb is None:
        if not os.path.exists(XLSX):
            raise FileNotFoundError(f"{XLSX} missing — download from {XLSX_URL}")
        _wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    return _wb


def block(sheet: str, row0: int, nrow: int, col0: int, ncol: int) -> np.ndarray:
    """1-based (row0, col0) top-left corner; returns float array (nrow, ncol), NaN where empty/non-numeric."""
    ws = workbook()[sheet]
    out = np.full((nrow, ncol), np.nan)
    for i, row in enumerate(ws.iter_rows(min_row=row0, max_row=row0 + nrow - 1, min_col=col0,
                                          max_col=col0 + ncol - 1, values_only=True)):
        for j, v in enumerate(row):
            try:
                out[i, j] = float(v)
            except (TypeError, ValueError):
                pass
    return out


def published_pxp():
    """(53, 3) protected exceedance probabilities [null, 2B, bifurcation] and their window centres (ms)."""
    px = block("Figure 3 Panel E", 2, 3, 1, 53).T
    t = -285.0 + 30.0 * np.arange(53)
    return t, px


def published_behaviour_sd():
    """(20, 5) SD of audibility, levels 2..6, on the 0-10 scale (the sheet is in %)."""
    return block("Figure 1 Panel D", 2, 20, 17, 6)[:, 1:] / 10.0


def published_neural_sd_active():
    """(20, 8, 5) variability profile of the projected activity, 8 article windows x levels 2..6."""
    return block("Figure 2 Panel C rightmost col", 26, 20, 1, 40).reshape(20, 8, 5)


def published_passive_sd():
    """(20, 6, 6) passive variability profile, 6 windows (0-100 .. 500-600 ms) x 6 columns."""
    return block("Supplementary Figure 4", 25, 20, 1, 36).reshape(20, 6, 6)


def inspect():
    wb = workbook()
    print("sheets:", wb.sheetnames)
    for name, rows in [("Figure 3 Panel E", (1, 6)), ("Figure 1 Panel D", (1, 3)),
                       ("Figure 2 Panel C rightmost col", (1, 27)), ("Supplementary Figure 4", (1, 26))]:
        ws = wb[name]
        print(f"===== {name}")
        for r in ws.iter_rows(min_row=rows[0], max_row=rows[1], values_only=True):
            cells = [(j + 1, v) for j, v in enumerate(r) if v is not None]
            strings = [(j, v) for j, v in cells if isinstance(v, str)]
            nums = [(j, v) for j, v in cells if not isinstance(v, str)]
            print("  strings:", strings[:12], "| numeric cells:", len(nums), "first:", nums[:3])
    # non-numeric label cells anywhere in the used blocks' neighbourhoods
    for name in ["Figure 2 Panel C rightmost col", "Supplementary Figure 4", "Figure 1 Panel D"]:
        ws = wb[name]
        labels = []
        for r_i, r in enumerate(ws.iter_rows(values_only=True), start=1):
            for j, v in enumerate(r, start=1):
                if isinstance(v, str):
                    labels.append((r_i, j, v[:40]))
        print(f"----- {name}: all string cells ({len(labels)}):", labels[:60])


if __name__ == "__main__":
    if "--inspect" in sys.argv:
        inspect()
